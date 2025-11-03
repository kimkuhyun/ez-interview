"""업로드된 문서에서 텍스트와 미리보기를 생성하는 유틸리티들."""

from __future__ import annotations

import base64
import html
import io
from pathlib import Path
from typing import Dict, List, Optional, TypedDict

import fitz  # PyMuPDF
import mammoth  # DOCX to HTML
import pypdfium2 as pdfium
from docx import Document
from fastapi import UploadFile
from PIL import Image

from .text_extract import extract_text_any


class PreviewPayload(TypedDict):
    pdf_images: List[str]
    docx_url: Optional[str]
    html_preview: Optional[str]



async def read_upload_bytes(file: UploadFile) -> bytes:
    return await file.read()

def extract_text_for_state(filename: str, file_bytes: bytes) -> str:
    return extract_text_any(filename, file_bytes)

def detect_extension(filename: str) -> str:
    """파일 확장자(소문자)를 반환합니다."""
    return Path(filename).suffix.lower().lstrip(".")


def render_pdf_to_b64_images(file_bytes: bytes, max_pages: int = 5, dpi: int = 300) -> List[str]:
    """PDF를 PNG Base64 이미지 리스트로 변환합니다 (폴백 포함)."""
    pages: List[str] = []
    
    try:
        # PyPdfium2로 먼저 시도
        document = pdfium.PdfDocument(io.BytesIO(file_bytes))
        scale = dpi / 72.0  # PDF 기본 DPI는 72
        
        for page in document[:max_pages]:
            bitmap = page.render(scale=scale)
            buffer = io.BytesIO()
            bitmap.to_pil().save(buffer, format="PNG", optimize=True)
            pages.append(base64.b64encode(buffer.getvalue()).decode("ascii"))
            
    except Exception as e:
        try:
            # PyMuPDF로 폴백
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            try:
                zoom = dpi / 72.0
                transform = fitz.Matrix(zoom, zoom)
                
                for index, page in enumerate(doc):
                    pixmap = page.get_pixmap(matrix=transform)
                    pages.append(base64.b64encode(pixmap.tobytes("png")).decode("ascii"))
                    if index + 1 >= max_pages:
                        break
            finally:
                doc.close()
        except Exception as e2:
            raise ValueError(f"PDF 렌더링 실패: {str(e2)}") from e2
            
    if not pages:
        raise ValueError("PDF에서 이미지를 추출할 수 없습니다")
        
    return pages


def _convert_table_to_html(table: Document.tables) -> List[str]:
    """DOCX 표를 HTML 스니펫으로 변환합니다."""
    snippets: List[str] = ["<table class='border-collapse border'>"]
    
    for row in table.rows:
        snippets.append("<tr>")
        for cell in row.cells:
            text = html.escape(cell.text.strip())
            snippets.append(f"<td class='border p-2'>{text}</td>")
        snippets.append("</tr>")
        
    snippets.append("</table>")
    return snippets

def render_docx_to_html(file_bytes: bytes, max_paragraphs: int = 200) -> str:
    """DOCX를 HTML로 변환하여 반환합니다 (폴백 포함)."""
    try:
        # Mammoth로 먼저 시도
        result = mammoth.convert_to_html(io.BytesIO(file_bytes))
        return result.value
        
    except Exception:
        try:
            # python-docx로 폴백
            document = Document(io.BytesIO(file_bytes))
            snippets: List[str] = []
            
            # 표 변환
            for table in document.tables:
                snippets.extend(_convert_table_to_html(table))
                snippets.append("<br>")
            
            # 문단 변환
            for index, paragraph in enumerate(document.paragraphs):
                if index >= max_paragraphs:
                    snippets.append("<p>... (나머지 내용 생략)</p>")
                    break
                    
                text = paragraph.text.strip()
                if text:
                    snippets.append(f"<p class='my-2'>{html.escape(text)}</p>")
                    
            return "\n".join(snippets) if snippets else '<p class="text-gray-500 text-sm">내용이 없습니다.</p>'
            
        except Exception as e:
            raise ValueError(f"DOCX HTML 변환 실패: {str(e)}") from e


def build_preview_payload(filename: str, file_bytes: bytes) -> PreviewPayload:
    """미리보기(이미지 또는 HTML) 데이터를 생성하여 반환합니다."""
    ext = detect_extension(filename)
    
    preview: PreviewPayload = {
        "pdf_images": [],
        "docx_url": None,
        "html_preview": None,
    }
    
    if ext == "pdf":
        preview["pdf_images"] = render_pdf_to_b64_images(file_bytes)
    elif ext == "docx":
        preview["html_preview"] = render_docx_to_html(file_bytes)
    else:
        raise ValueError(f"지원하지 않는 파일 형식입니다: {ext}")
        
    return preview


def process_bytes(filename: str, file_bytes: bytes) -> tuple[str, PreviewPayload]:
    """파일 바이트를 받아 텍스트와 미리보기 페이로드를 생성하여 반환합니다."""
    text = extract_text_any(filename, file_bytes)
    preview = build_preview_payload(filename, file_bytes)
    return text, preview
