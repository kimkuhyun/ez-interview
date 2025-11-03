"""문서 텍스트 추출 유틸리티."""

from io import BytesIO
from typing import List
import fitz  # PyMuPDF
from docx import Document
from docx.table import Table, _Row
from docx.text.paragraph import Paragraph


def _extract_text_pdf(file_bytes: bytes) -> str:
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    try:
        parts = []
        for page in doc:
            text = page.get_text("text").strip()
            if text:
                parts.append(text)
        return "\n".join(parts)
    finally:
        doc.close()

def _extract_table_text(table: Table) -> List[str]:
    texts = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
        if cells:
            texts.append("\t".join(cells))
    return texts

def _extract_text_docx(file_bytes: bytes) -> str:
    try:
        doc = Document(BytesIO(file_bytes))
    except Exception as e:
        raise ValueError(f"DOCX 파일 파싱 실패: {str(e)}") from e
        
    parts: List[str] = []
    
    # 표 내용 수집
    for table in doc.tables:
        parts.extend(_extract_table_text(table))
        
    # 문단 수집
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)
            
    return "\n".join(parts)

def extract_text_any(filename: str, file_bytes: bytes) -> str:
    name = filename.lower()
    
    if name.endswith(".pdf"):
        try:
            return _extract_text_pdf(file_bytes)
        except Exception as e:
            raise ValueError(f"PDF 텍스트 추출 실패: {str(e)}") from e
            
    if name.endswith(".docx"):
        return _extract_text_docx(file_bytes)
        
    raise ValueError(f"지원하지 않는 파일 형식: {filename}")